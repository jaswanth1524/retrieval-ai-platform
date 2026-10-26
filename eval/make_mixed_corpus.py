"""Write ``eval/corpus_mixed``: the handbook corpus with six documents in other formats.

The retrieval eval's corpora are Markdown, so the PDF, DOCX, HTML, ODT, PPTX and XLSX
parsers never reached a measured run. This renders the same text into each format —
travel → PDF (a page per section), expense → DOCX, leave → HTML, IT equipment → ODT,
security → PPTX (a slide per section), procurement → XLSX (a row per paragraph) — and
copies the rest, so ``eval/datasets/mixed_eval.jsonl`` (the handbook questions, renamed)
scores the parsers against the Markdown baseline.

    uv run python -m eval.make_mixed_corpus

Regenerate it only when the handbook corpus changes; Office formats embed timestamps,
so the bytes differ from run to run.
"""

from __future__ import annotations

import html
import io
import json
import shutil
import sys
import zipfile
from pathlib import Path

EVAL = Path(__file__).resolve().parent
SOURCE = EVAL / "corpus_handbook"
TARGET = EVAL / "corpus_mixed"

# Handbook document -> the format it is rendered in.
RENDERED = {
    "travel_policy.md": ".pdf",
    "expense_policy.md": ".docx",
    "leave_policy.md": ".html",
    "it_equipment_policy.md": ".odt",
    "security_policy.md": ".pptx",
    "procurement_policy.md": ".xlsx",
}

Section = tuple[str, list[str]]


def read_markdown(path: Path) -> tuple[str, list[Section]]:
    """A handbook file's title and its ``## `` sections (heading, paragraphs)."""

    title = path.stem.replace("_", " ").capitalize()
    sections: list[Section] = []
    paragraph: list[str] = []

    def end_paragraph() -> None:
        if paragraph and sections:
            sections[-1][1].append(" ".join(paragraph))
        paragraph.clear()

    for line in path.read_text(encoding="utf-8").splitlines():
        stripped = line.strip()
        if stripped.startswith("# "):
            title = stripped[2:].strip()
        elif stripped.startswith("## "):
            end_paragraph()
            sections.append((stripped[3:].strip(), []))
        elif not stripped:
            end_paragraph()
        else:
            paragraph.append(stripped)
    end_paragraph()
    return title, sections


def _wrap(text: str, width: int = 90) -> list[str]:
    lines: list[str] = []
    current = ""
    for word in text.split():
        if current and len(current) + 1 + len(word) > width:
            lines.append(current)
            current = word
        else:
            current = f"{current} {word}" if current else word
    if current:
        lines.append(current)
    return lines


def render_pdf(title: str, sections: list[Section]) -> bytes:
    """A page per section, Helvetica text, lines wrapped at 90 characters."""

    pages = []
    for index, (heading, paragraphs) in enumerate(sections):
        lines = [title, ""] if index == 0 else []
        lines.append(heading)
        for paragraph in paragraphs:
            lines.extend(_wrap(paragraph))
        pages.append(lines)

    def escape(line: str) -> str:
        return line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    objects: list[bytes] = [b"<< /Type /Catalog /Pages 2 0 R >>"]
    kids = " ".join(f"{3 + 2 * index} 0 R" for index in range(len(pages)))
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    font_id = 3 + 2 * len(pages)
    for index, lines in enumerate(pages):
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> "
            f"/Contents {4 + 2 * index} 0 R >>".encode()
        )
        shown = " 0 -14 Td ".join(f"({escape(line)}) Tj" for line in lines)
        stream = f"BT /F1 11 Tf 56 740 Td {shown} ET".encode("latin-1")
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    objects.append(
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>"
    )
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return out.getvalue()


def render_docx(title: str, sections: list[Section]) -> bytes:
    from docx import Document

    document = Document()
    document.add_heading(title, level=1)
    for heading, paragraphs in sections:
        document.add_heading(heading, level=2)
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def render_html(title: str, sections: list[Section]) -> bytes:
    body = "".join(
        f"<section><h2>{html.escape(heading)}</h2>"
        + "".join(f"<p>{html.escape(paragraph)}</p>" for paragraph in paragraphs)
        + "</section>"
        for heading, paragraphs in sections
    )
    page = (
        "<!doctype html><html lang='en'><head><meta charset='utf-8'>"
        f"<title>{html.escape(title)}</title></head><body>"
        "<nav><a href='/'>Intranet</a> · <a href='/policies'>Policies</a></nav>"
        f"<main><h1>{html.escape(title)}</h1>{body}</main>"
        "<footer>Internal use only.</footer></body></html>"
    )
    return page.encode()


def render_odt(title: str, sections: list[Section]) -> bytes:
    def h(text: str, level: int) -> str:
        return f'<text:h text:outline-level="{level}">{html.escape(text)}</text:h>'

    body = h(title, 1) + "".join(
        h(heading, 2) + "".join(f"<text:p>{html.escape(p)}</text:p>" for p in paragraphs)
        for heading, paragraphs in sections
    )
    content = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        "<office:document-content"
        ' xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"'
        ' xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"'
        ' office:version="1.3">'
        f"<office:body><office:text>{body}</office:text></office:body>"
        "</office:document-content>"
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        archive.writestr("content.xml", content)
    return buffer.getvalue()


def render_pptx(title: str, sections: list[Section]) -> bytes:
    from pptx import Presentation

    deck = Presentation()
    cover = deck.slides.add_slide(deck.slide_layouts[0])
    cover.shapes.title.text = title
    for heading, paragraphs in sections:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = heading
        frame = slide.placeholders[1].text_frame
        frame.text = paragraphs[0] if paragraphs else ""
        for paragraph in paragraphs[1:]:
            frame.add_paragraph().text = paragraph
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def render_xlsx(title: str, sections: list[Section]) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Policy"
    sheet.append(["Section", "Text"])
    for heading, paragraphs in sections:
        for paragraph in paragraphs:
            sheet.append([heading, paragraph])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


RENDERERS = {
    ".pdf": render_pdf,
    ".docx": render_docx,
    ".html": render_html,
    ".odt": render_odt,
    ".pptx": render_pptx,
    ".xlsx": render_xlsx,
}


def renamed(name: str) -> str:
    """A handbook file's name in the mixed corpus."""

    suffix = RENDERED.get(name)
    return name if suffix is None else Path(name).with_suffix(suffix).name


def build_dataset(source: Path, target: Path) -> int:
    """The handbook questions with their files' names in the mixed corpus."""

    rows = []
    for line in source.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        row = json.loads(line)
        if "relevant_filenames" in row:
            row["relevant_filenames"] = [renamed(name) for name in row["relevant_filenames"]]
        rows.append(json.dumps(row, ensure_ascii=False))
    target.write_text("\n".join(rows) + "\n", encoding="utf-8")
    return len(rows)


def build(target: Path = TARGET) -> list[str]:
    if target.exists():
        shutil.rmtree(target)
    target.mkdir(parents=True)
    written = []
    for path in sorted(SOURCE.glob("*.md")):
        name = renamed(path.name)
        if name == path.name:
            shutil.copyfile(path, target / name)
        else:
            title, sections = read_markdown(path)
            (target / name).write_bytes(RENDERERS[Path(name).suffix](title, sections))
        written.append(name)
    return written


def main() -> int:
    for name in build():
        print(name)
    count = build_dataset(
        EVAL / "datasets" / "handbook_eval.jsonl", EVAL / "datasets" / "mixed_eval.jsonl"
    )
    print(f"{count} questions -> eval/datasets/mixed_eval.jsonl")
    return 0


if __name__ == "__main__":
    sys.exit(main())
