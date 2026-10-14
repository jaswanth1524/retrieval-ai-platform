"""PPTX parsing: each slide a page, its title the section, speaker notes included."""

from __future__ import annotations

from collections.abc import Iterable, Iterator
from io import BytesIO
from typing import Any

from api.parsers.common import (
    ZIP_MAX_EXPANDED_BYTES,
    DocumentParseError,
    DocumentSection,
    EmptyDocumentError,
    cap_section,
    check_zip_expanded_size,
    normalize_text,
    refuse_zip_xml_dtd,
    render_row,
)

# python-pptx parses each slide's XML in memory; see ZIP_MAX_EXPANDED_BYTES.
PPTX_MAX_EXPANDED_BYTES = ZIP_MAX_EXPANDED_BYTES

# Group shapes nest; real decks go a few levels deep at most.
_MAX_GROUP_DEPTH = 16


def parse_pptx_document(filename: str, content: bytes) -> list[DocumentSection]:
    """Extract each slide's text: text boxes and placeholders in shape order, tables as
    "header: value; ..." rows, then the speaker notes.

    A slide is a page (its number in the deck) and its title is the section label
    ("Slide n" without one), so a citation reads as the slide a presenter would show.
    Charts, pictures and media are skipped.
    """

    check_zip_expanded_size(filename, content, "PPTX", PPTX_MAX_EXPANDED_BYTES)
    refuse_zip_xml_dtd(filename, content, "PPTX")

    from pptx import Presentation

    try:
        presentation = Presentation(BytesIO(content))
        sections: list[DocumentSection] = []
        for number, slide in enumerate(presentation.slides, start=1):
            title_shape = slide.shapes.title
            title = normalize_text(title_shape.text_frame.text) if title_shape is not None else ""
            # By id: python-pptx hands out a new proxy object per lookup. The title is
            # the section label (and in every chunk's embedded prefix), not body text.
            title_id = title_shape.shape_id if title_shape is not None else None
            parts = [
                text
                for shape in _shapes(slide.shapes, 0)
                if shape.shape_id != title_id and (text := _shape_text(shape))
            ]
            if slide.has_notes_slide:
                notes = normalize_text(slide.notes_slide.notes_text_frame.text or "")
                if notes:
                    parts.append(f"Notes: {notes}")
            text = normalize_text("\n".join(parts))
            if not text and not title:
                continue
            sections.append(
                DocumentSection(
                    filename=filename,
                    page=number,
                    section=cap_section(title or f"Slide {number}"),
                    # A title-only slide (a section divider) still says something.
                    text=text or title,
                    # Never merged into the next slide: short slides are the norm.
                    boundary="isolated",
                )
            )
    except DocumentParseError:
        raise
    except Exception as exc:  # python-pptx and its XML layer raise many types.
        raise DocumentParseError(f"Could not parse PPTX '{filename}'.") from exc
    if not sections:
        raise EmptyDocumentError(f"PPTX '{filename}' has no slide text.")
    return sections


def _shapes(shapes: Iterable[Any], depth: int) -> Iterator[Any]:
    """Every shape, group members included, in the order they are drawn."""

    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            if depth < _MAX_GROUP_DEPTH:
                yield from _shapes(shape.shapes, depth + 1)
            continue
        yield shape


def _shape_text(shape: Any) -> str:
    if getattr(shape, "has_table", False) and shape.has_table:
        rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
        if not rows:
            return ""
        header = [cell.strip() for cell in rows[0]]
        rendered = [render_row(header, row) for row in rows[1:]]
        lines = [line for line in rendered if line] or ["; ".join(cell for cell in header if cell)]
        return "\n".join(lines)
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        return normalize_text(shape.text_frame.text)
    return ""
