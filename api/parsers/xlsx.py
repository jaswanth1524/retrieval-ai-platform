"""XLSX parsing: each sheet a page, its rows rendered like a CSV's."""

from __future__ import annotations

import datetime
from io import BytesIO

from api.parsers.common import (
    ZIP_MAX_EXPANDED_BYTES,
    DocumentParseError,
    DocumentSection,
    EmptyDocumentError,
    cap_section,
    check_zip_expanded_size,
    refuse_zip_xml_dtd,
    row_groups,
)
from api.settings import AppSettings

# openpyxl reads sheet XML as it goes, but shared strings are held in memory; see
# ZIP_MAX_EXPANDED_BYTES.
XLSX_MAX_EXPANDED_BYTES = ZIP_MAX_EXPANDED_BYTES


def parse_xlsx_document(
    filename: str, content: bytes, settings: AppSettings | None = None
) -> list[DocumentSection]:
    """Render each sheet's rows as "header: value; ..." text, in row-range sections.

    A sheet is a page (its position in the workbook), so a citation names the sheet by
    page number and by its section label, ``"<sheet> · Rows 4-53"`` — spreadsheet row
    numbers, blank rows included, as with CSV. The first non-blank row of a sheet is its
    header. Cells hold their last computed values (``data_only``); formulas are not
    evaluated. Charts and images are skipped.
    """

    limits = settings if settings is not None else AppSettings.model_construct()
    rows_per_section = int(limits.csv_rows_per_section)
    max_cells = int(limits.max_spreadsheet_cells)
    check_zip_expanded_size(filename, content, "XLSX", XLSX_MAX_EXPANDED_BYTES)
    refuse_zip_xml_dtd(filename, content, "XLSX")

    from openpyxl import load_workbook

    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several reader-specific exceptions.
        raise DocumentParseError(f"Could not parse XLSX '{filename}'.") from exc

    sections: list[DocumentSection] = []
    cells = 0
    try:
        for page, sheet in enumerate(workbook.worksheets, start=1):
            # The declared dimension is only a hint; trusting it padded every row to
            # its width (a sheet claiming A1:XFD1048576 is 17 billion cells).
            sheet.reset_dimensions()
            records: list[tuple[int, list[str]]] = []
            for number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                cells += len(row) + 1
                if cells > max_cells:
                    raise DocumentParseError(
                        f"XLSX '{filename}' has more than {max_cells} cells "
                        "(MAX_SPREADSHEET_CELLS). Split it into smaller files."
                    )
                values = [_cell_text(value) for value in row]
                if any(value.strip() for value in values):
                    records.append((number, values))
            sections.extend(_sheet_sections(filename, page, sheet.title, records, rows_per_section))
    except DocumentParseError:
        raise
    except Exception as exc:
        raise DocumentParseError(f"Could not parse XLSX '{filename}'.") from exc
    finally:
        workbook.close()
    if not sections:
        raise EmptyDocumentError(f"XLSX '{filename}' has no cell values.")
    return sections


def _sheet_sections(
    filename: str,
    page: int,
    title: str,
    records: list[tuple[int, list[str]]],
    rows_per_section: int,
) -> list[DocumentSection]:
    if not records:
        return []
    header = [cell.strip() for cell in records[0][1]]
    data_rows = records[1:]
    if not data_rows:
        return [
            DocumentSection(
                filename=filename,
                page=page,
                section=cap_section(f"{title} · Header"),
                text="; ".join(cell for cell in header if cell),
                boundary="isolated",
            )
        ]
    sections: list[DocumentSection] = []
    for first, last, rendered in row_groups(header, data_rows, rows_per_section):
        label = f"{title} · Row {first}" if first == last else f"{title} · Rows {first}-{last}"
        sections.append(
            DocumentSection(
                filename=filename,
                page=page,
                section=cap_section(label),
                text="\n".join(line for _, line in rendered),
                boundary="isolated",
                # Chunks keep rows whole and are labelled with the rows they hold.
                line_rows=tuple(number for number, _ in rendered),
                row_label_prefix=f"{title[:80]} · ",
            )
        )
    return sections


def _cell_text(value: object) -> str:
    """A cell value as it reads in the sheet: 3 rather than 3.0, dates without a
    midnight time."""

    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime.datetime):
        if value.time() == datetime.time(0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, datetime.date | datetime.time):
        return value.isoformat()
    return str(value)
