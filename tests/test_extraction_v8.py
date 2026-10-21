"""Text extraction fixes of CHUNKER_VERSION 8: shared clean-up, broken ligature maps,
words broken over line ends, book running heads, stamped scans, DOCX content python-docx
skipped, binary files renamed to .txt."""

from __future__ import annotations

import io
import logging
from typing import Any

import pytest

from api.documents import parse_document_bytes
from api.parsers import pdf as pdf_parser
from api.parsers.common import DocumentParseError, DocumentSection, clean_text
from api.parsers.pdf import rejoin_hyphenated, strip_running_lines
from tests.samples import pdf_bytes


def test_clean_text_writes_out_ligatures_and_drops_word_splitting_characters() -> None:
    assert clean_text("ﬁnance ﬂoor") == "finance floor"
    assert clean_text("infor­mation zero​width joi⁠ner﻿") == ("information zerowidth joiner")
    # The joiner and non-joiner are part of words in Persian and Indic scripts.
    assert clean_text("می‌خوام") == "می‌خوام"
    # Not NFKC: superscripts and fractions are left as written.
    assert clean_text("m² ½") == "m² ½"


def test_every_parsed_document_and_question_is_cleaned() -> None:
    sections = parse_document_bytes("notes.txt", "The ﬁscal year starts in April.".encode())
    assert sections[0].text == "The fiscal year starts in April."
    with pytest.raises(Exception, match="no extractable text"):
        parse_document_bytes("blank.txt", "​­".encode())

    from api.embeddings import EmbeddedText
    from api.retrieval import embed_query

    seen: list[str] = []

    class Recording:
        def embed_texts(self, texts: list[str]) -> list[EmbeddedText]:
            seen.extend(texts)
            return [EmbeddedText(dense=[1.0], sparse=None)]  # type: ignore[arg-type]

    embed_query("ﬁscal year", Recording())  # type: ignore[arg-type]
    assert seen == ["fiscal year"]


def test_a_font_whose_map_contradicts_its_ligature_glyphs_is_read_right() -> None:
    content = pdf_bytes(
        ['The speci!c rule: book the !rst "ight.', 'Desks on the "oor are \\(re\\)used.'] * 2,
        ligature_fault=True,
    )
    text = "\n".join(section.text for section in parse_document_bytes("book.pdf", content))
    assert "specific rule: book the first flight." in text
    assert "Desks on the floor" in text
    assert "!" not in text and '"' not in text


def test_a_correct_ligature_map_is_left_as_it_was() -> None:
    sections = parse_document_bytes("plain.pdf", pdf_bytes(["Wow! A plain page."]))
    assert sections[0].text == "Wow! A plain page."


def test_words_broken_over_a_line_end_are_made_whole() -> None:
    pages = [
        "Every pur-\nchase needs a quote. Ask the ded-\nicated team.",
        "A cost-\neffective plan: the cost is low and effective.",
        "Call the help-\ndesk. The help-desk is open. Pre-\nCOVID rules ended.",
        "Install the up-\ndate; an update is due.",
    ]
    assert rejoin_hyphenated(pages) == [
        "Every purchase needs a quote. Ask the dedicated team.",
        "A cost-effective plan: the cost is low and effective.",
        "Call the help-desk. The help-desk is open. Pre-COVID rules ended.",
        "Install the update; an update is due.",
    ]


def test_a_book_running_head_with_its_page_number_is_removed() -> None:
    # Left pages: "<n> Chapter 1: Prerequisites"; right pages: "1.2 About Your App <n>".
    pages = [
        (f"{25 + n} Chapter 1: Prerequisites" if n % 2 else f"1.2 About Your App {25 + n}")
        + f"\nBody text of page {n}.\nMore of page {n}."
        for n in range(1, 7)
    ]
    assert strip_running_lines(pages) == [
        f"Body text of page {n}.\nMore of page {n}." for n in range(1, 7)
    ]
    # A one-word heading numbered with the page is content ("Step 3" on page 3).
    steps = [f"Step {n}\nDo thing {n} first.\nThen check part {n}." for n in range(1, 6)]
    assert strip_running_lines(steps) == steps


def test_a_scanned_page_with_only_a_stamp_is_ocred(monkeypatch: pytest.MonkeyPatch) -> None:
    content = pdf_bytes(
        ["Contract terms in full text on this page.", "ACME-000123"],
        image_pages=frozenset({2}),
    )
    asked: list[list[int]] = []

    def fake_ocr(filename: str, data: bytes, pages: list[int]) -> list[DocumentSection]:
        asked.append(pages)
        return [DocumentSection(filename, 2, "Scan", "Signed by both parties on 2 May.")]

    monkeypatch.setattr(pdf_parser, "_ocr_pdf_pages", fake_ocr)
    sections = parse_document_bytes("contract.pdf", content)
    assert asked == [[2]]
    assert [s.text for s in sections] == [
        "Contract terms in full text on this page.",
        "Signed by both parties on 2 May.",
    ]

    # OCR not installed: the stamp's own text is kept rather than the page lost.
    def no_ocr(filename: str, data: bytes, pages: list[int]) -> list[DocumentSection]:
        raise pdf_parser.EmptyDocumentError("no OCR")

    monkeypatch.setattr(pdf_parser, "_ocr_pdf_pages", no_ocr)
    sections = parse_document_bytes("contract.pdf", content)
    assert [s.text for s in sections][-1] == "ACME-000123"
    # A short page without a scan on it is just a short page.
    plain = parse_document_bytes("short.pdf", pdf_bytes(["Long text page.", "Annex B"]))
    assert [s.text for s in plain] == ["Long text page.", "Annex B"]


def test_a_pipe_inside_a_table_cell_does_not_split_it() -> None:
    table = pdf_parser._table_to_markdown([["Plan", "Notes"], ["Basic", "a|b"]])
    assert "| Basic | a\\|b |" in table


def test_a_binary_file_renamed_to_txt_is_refused() -> None:
    with pytest.raises(DocumentParseError, match="binary"):
        parse_document_bytes("photo.txt", bytes(range(256)) * 20)
    # Tabs and form feeds are text.
    assert parse_document_bytes("ok.txt", b"a\tb\fc\nd")


def _docx(build: Any) -> bytes:
    from docx import Document

    document = Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_text_in_content_controls_tracked_insertions_and_fields_is_read() -> None:
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    def build(document: Any) -> None:
        document.add_heading("Leave", level=1)
        paragraph = document.add_paragraph("Staff get ")
        paragraph._p.append(
            parse_xml(
                f'<w:ins {nsdecls("w")} w:id="1" w:author="a"><w:r><w:t>twenty days</w:t>'
                "</w:r></w:ins>"
            )
        )
        paragraph._p.append(
            parse_xml(
                f'<w:del {nsdecls("w")} w:id="2" w:author="a"><w:r><w:delText>ten days'
                "</w:delText></w:r></w:del>"
            )
        )
        paragraph.add_run(" of leave.")
        body = document.element.body
        body.insert(
            len(body) - 1,
            parse_xml(
                f"<w:sdt {nsdecls('w')}><w:sdtContent><w:p><w:r><w:t>Carry-over ends in "
                "March.</w:t></w:r></w:p></w:sdtContent></w:sdt>"
            ),
        )
        see = document.add_paragraph("See ")
        see._p.append(
            parse_xml(
                f'<w:fldSimple {nsdecls("w")} w:instr="REF x"><w:r><w:t>section 4</w:t>'
                "</w:r></w:fldSimple>"
            )
        )

    text = parse_document_bytes("leave.docx", _docx(build))[0].text
    assert text == (
        "Leave\nStaff get twenty days of leave.\nCarry-over ends in March.\nSee section 4"
    )


def test_docx_tables_read_each_cell_once_and_label_their_rows() -> None:
    def build(document: Any) -> None:
        small = document.add_table(rows=2, cols=3)
        for index, value in enumerate(["Grade", "Days", "Note"]):
            small.cell(0, index).text = value
        small.cell(1, 0).merge(small.cell(1, 1)).text = "Merged"
        small.cell(1, 2).text = "x"
        big = document.add_table(rows=3, cols=2)
        for row, values in enumerate([("Grade", "Days"), ("A", "20"), ("B", "25")]):
            for column, value in enumerate(values):
                big.cell(row, column).text = value

    text = parse_document_bytes("grades.docx", _docx(build))[0].text
    assert (
        text
        == "Grade | Days | Note\nMerged | x\nGrade | Days\nGrade: A; Days: 20\nGrade: B; Days: 25"
    )


def test_parser_warnings_are_logged_once_per_document(caplog: pytest.LogCaptureFixture) -> None:
    from api.parse_isolation import summarized_parser_warnings

    caplog.set_level(logging.WARNING)
    with summarized_parser_warnings("odd.pdf"):
        for index in range(5):
            logging.getLogger("pypdf").warning("font problem %d " + "x" * 500, index)
    messages = [record.getMessage() for record in caplog.records]
    assert len(messages) == 1
    assert "odd.pdf" in messages[0] and "5 problem(s)" in messages[0]
    assert len(messages[0]) < 400


def test_a_running_head_that_changes_by_subsection_loses_only_its_page_number() -> None:
    # Left pages repeat the chapter (and prove the page-number offset, 25); right pages
    # carry the subsection, which changes, so only its number can go.
    heads = {
        2: "1.2 About Your App 27",
        4: "1.2.2 Online versus Offline 29",
        6: "1.3 Hardware Choices 31",
    }
    pages = [
        (heads[n] if n in heads else f"{25 + n} Chapter 1: Prerequisites")
        + f"\nBody text of page {n}.\nMore of page {n}."
        for n in range(1, 7)
    ]
    cleaned = strip_running_lines(pages)
    assert cleaned[0] == "Body text of page 1.\nMore of page 1."
    assert cleaned[1] == "1.2 About Your App\nBody text of page 2.\nMore of page 2."
    assert cleaned[3].startswith("1.2.2 Online versus Offline\n")
    # Without a proven offset, a trailing number is left alone.
    alone = [f"Plan {n} Overview {n}\nText {n}.\nEnd of {n} here." for n in range(1, 5)]
    assert strip_running_lines(alone) == alone


def test_a_budget_strips_every_neighbour_before_dropping_a_passage() -> None:
    # A dropped passage's text was left out of its neighbour's expansion (it had its own
    # block); the budget takes all neighbour text first, so nothing of it is lost that the
    # budget would have kept.
    from api.prompting import (
        build_grounded_messages,
        estimate_prompt_tokens,
        fit_prompt_to_context_window,
    )
    from api.reranking import RerankedChunk
    from tests.factories import make_test_settings

    def passage(n: int) -> RerankedChunk:
        return RerankedChunk(
            point_id=f"p{n}",
            filename="a.md",
            page=1,
            section="S",
            chunk_id=f"c{n}",
            text=f"passage {n} " + "word " * 60,
            retrieval_score=0.5,
            rerank_score=1 - n / 10,
            chunk_ordinal=n,
            expanded_text="neighbour " * 200,
        )

    chunks = [passage(1), passage(2), passage(3)]
    two_bare = [RerankedChunk(**{**c.__dict__, "expanded_text": None}) for c in chunks[:2]]
    budget = estimate_prompt_tokens(build_grounded_messages("q", two_bare))
    fitted = fit_prompt_to_context_window(
        "q", chunks, None, make_test_settings(llm_provider="ollama", prompt_token_budget=budget)
    )
    assert [c.chunk_id for c in fitted.context_chunks] == ["c1", "c2"]
    assert all(c.expanded_text is None for c in fitted.context_chunks)
