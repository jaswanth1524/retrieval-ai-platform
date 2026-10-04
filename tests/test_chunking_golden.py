"""Golden snapshot of parse + chunk output across every supported format.

Pins the exact chunks (id, page, section, text hash) the parsers and the chunker produce
for a fixed set of documents, so moving parser code between modules can't change what
gets indexed without this failing. A deliberate chunking change regenerates it:

    DOCRAG_UPDATE_GOLDEN=1 uv run pytest tests/test_chunking_golden.py

and, if the change alters chunk boundaries, bumps ``CHUNKER_VERSION``.
"""

from __future__ import annotations

import hashlib
import io
import json
import os
from pathlib import Path
from typing import Any

import pytest
from docx import Document

from api.chunking import HeuristicTokenCounter
from api.documents import chunk_sections, parse_document_bytes
from api.settings import AppSettings
from tests.samples import pdf_bytes

GOLDEN_PATH = Path(__file__).parent / "golden" / "chunks.json"
REPO_ROOT = Path(__file__).resolve().parent.parent


def _settings() -> AppSettings:
    return AppSettings(  # type: ignore[call-arg]
        _env_file=None, chunk_size_tokens=60, chunk_overlap_tokens=10
    )


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("Leave policy", level=1)
    document.add_paragraph(
        "Employees accrue twenty days of paid leave per year. Unused days carry over."
    )
    document.add_heading("Approvals", level=2)
    document.add_paragraph("Managers approve leave requests within two working days.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Type"
    table.cell(0, 1).text = "Days"
    table.cell(1, 0).text = "Annual"
    table.cell(1, 1).text = "20"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _documents() -> dict[str, bytes]:
    documents: dict[str, bytes] = {}
    for corpus in ("eval/corpus", "eval/corpus_handbook"):
        for path in sorted((REPO_ROOT / corpus).glob("*.md")):
            documents[f"{corpus.split('/')[-1]}__{path.name}"] = path.read_bytes()
    documents["sample.txt"] = (
        b"Plain notes about the release.\n\nThe second paragraph mentions qdrant and BM25."
    )
    documents["cjk.txt"] = "第一句话。第二句话！第三句话？最后一句。".encode()
    documents["page.html"] = (
        b"<html><head><script>ignored()</script></head><body><h1>Guide</h1>"
        b"<p>Install the stack with docker compose.</p><h2>Upgrade</h2>"
        b"<p>Pull the new image and restart.</p></body></html>"
    )
    documents["rows.csv"] = b"name,team\nada,platform\ngrace,search\nlin,ops\n"
    documents["policy.docx"] = _docx_bytes()
    documents["report.pdf"] = pdf_bytes(
        ["Quarterly report page one about retrieval quality.", "Page two covers latency."]
    )
    return documents


def _snapshot() -> dict[str, list[dict[str, Any]]]:
    settings = _settings()
    counter = HeuristicTokenCounter()
    snapshot: dict[str, list[dict[str, Any]]] = {}
    for name, content in _documents().items():
        upload_name = name.split("__")[-1]
        chunks = chunk_sections(
            parse_document_bytes(upload_name, content, settings), settings, counter
        )
        snapshot[name] = [
            {
                "chunk_id": chunk.chunk_id,
                "page": chunk.page,
                "section": chunk.section,
                "text_sha1": hashlib.sha1(chunk.text.encode()).hexdigest()[:16],
            }
            for chunk in chunks
        ]
    return snapshot


def test_parse_and_chunk_output_matches_the_golden_snapshot() -> None:
    actual = _snapshot()
    if os.environ.get("DOCRAG_UPDATE_GOLDEN") == "1":
        GOLDEN_PATH.parent.mkdir(exist_ok=True)
        GOLDEN_PATH.write_text(json.dumps(actual, indent=1, ensure_ascii=False) + "\n")
        pytest.skip("golden snapshot regenerated")
    expected = json.loads(GOLDEN_PATH.read_text())
    assert actual.keys() == expected.keys()
    for name in expected:
        assert actual[name] == expected[name], name
