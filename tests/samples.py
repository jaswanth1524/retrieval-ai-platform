"""Builders for small, real sample documents used across tests."""

from __future__ import annotations

import io


def pdf_bytes(
    pages: list[str],
    *,
    ligature_fault: bool = False,
    image_pages: frozenset[int] = frozenset(),
) -> bytes:
    """A minimal valid PDF of Helvetica text, one entry per page ("\n" between lines).

    ``ligature_fault``: the font's /Differences puts the fi and fl glyphs at codes 33
    and 34, while its /ToUnicode map says those codes are "!" and '"' (a real book's
    fault: "speci!c"). ``image_pages`` (1-based) also draw a large page-sized image.
    """

    objects: list[bytes] = []
    kids = " ".join(f"{3 + 2 * index} 0 R" for index in range(len(pages)))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    font_id = 3 + 2 * len(pages)
    cmap_id = font_id + 1
    image_id = font_id + 2
    for index, text in enumerate(pages):
        content_id = 4 + 2 * index
        image = f"/XObject << /Im1 {image_id} 0 R >> " if index + 1 in image_pages else ""
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font_id} 0 R >> {image}>> "
            f"/Contents {content_id} 0 R >>".encode()
        )
        # One line of text per "\n", each 14 points below the last.
        shown = " 0 -14 Td ".join(f"({line}) Tj" for line in text.split("\n"))
        paint = "q 612 0 0 792 0 0 cm /Im1 Do Q " if index + 1 in image_pages else ""
        stream = f"{paint}BT /F1 12 Tf 72 720 Td {shown} ET".encode()
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    if ligature_fault:
        objects.append(
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding << /Type "
            b"/Encoding /Differences [33 /fi 34 /fl] >> /ToUnicode %d 0 R >>" % cmap_id
        )
    else:
        objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    cmap = (
        b"/CIDInit /ProcSet findresource begin 12 dict begin begincmap /CMapName /X def "
        b"1 begincodespacerange <00> <FF> endcodespacerange 1 beginbfrange <20> <7E> <0020> "
        b"endbfrange endcmap CMapName currentdict /CMap defineresource pop end end"
    )
    objects.append(b"<< /Length %d >>\nstream\n" % len(cmap) + cmap + b"\nendstream")
    pixels = b"\x80" * 16
    objects.append(
        b"<< /Type /XObject /Subtype /Image /Width 1000 /Height 1000 /ColorSpace /DeviceGray "
        b"/BitsPerComponent 8 /Length %d >>\nstream\n" % len(pixels) + pixels + b"\nendstream"
    )

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return out.getvalue()


def epub_bytes(chapters: list[tuple[str, str]], *, spine_order: list[int] | None = None) -> bytes:
    """A minimal valid EPUB: one XHTML file per ``(title, body)`` chapter, read in
    ``spine_order`` (chapter indexes; default: as given)."""

    import zipfile

    order = spine_order if spine_order is not None else list(range(len(chapters)))
    manifest = "".join(
        f'<item id="c{index}" href="text/chapter%20{index}.xhtml" '
        'media-type="application/xhtml+xml"/>'
        for index in range(len(chapters))
    )
    spine = "".join(f'<itemref idref="c{index}"/>' for index in order)
    package = (
        '<?xml version="1.0" encoding="utf-8"?>'
        '<package xmlns="http://www.idpf.org/2007/opf" version="3.0">'
        f"<manifest>{manifest}"
        '<item id="css" href="style.css" media-type="text/css"/></manifest>'
        f"<spine>{spine}</spine></package>"
    )
    container = (
        '<?xml version="1.0"?>'
        '<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container" version="1.0">'
        '<rootfiles><rootfile full-path="OEBPS/content.opf" '
        'media-type="application/oebps-package+xml"/></rootfiles></container>'
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("mimetype", "application/epub+zip")
        archive.writestr("META-INF/container.xml", container)
        archive.writestr("OEBPS/content.opf", package)
        archive.writestr("OEBPS/style.css", "body { margin: 0 }")
        for index, (title, body) in enumerate(chapters):
            archive.writestr(
                f"OEBPS/text/chapter {index}.xhtml",
                '<?xml version="1.0" encoding="utf-8"?>'
                '<html xmlns="http://www.w3.org/1999/xhtml">'
                f"<head><title>{title}</title></head>"
                f"<body><h1>{title}</h1><p>{body}</p></body></html>",
            )
    return buffer.getvalue()


def eml_bytes(
    *, subject: str, body: str, html: str | None = None, attachment: bool = False
) -> bytes:
    """A small RFC 5322 email with a plain-text body, optionally an HTML alternative
    and a PDF attachment."""

    from email.message import EmailMessage

    message = EmailMessage()
    message["Subject"] = subject
    message["From"] = "Ada Lovelace <ada@example.com>"
    message["To"] = "team@example.com"
    message["Date"] = "Mon, 05 Oct 2026 09:30:00 +0000"
    message.set_content(body)
    if html is not None:
        message.add_alternative(html, subtype="html")
    if attachment:
        message.add_attachment(
            b"%PDF-1.4 not indexed", maintype="application", subtype="pdf", filename="a.pdf"
        )
    return bytes(message)


def xlsx_bytes(sheets: dict[str, list[list[object]]]) -> bytes:
    """A workbook with one sheet per entry, rows appended in order."""

    import io

    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets.items():
        sheet = workbook.create_sheet(title)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def pptx_bytes(slides: list[dict[str, object]]) -> bytes:
    """A deck: each slide may have a ``title``, ``body`` lines, a ``table`` (rows) and
    ``notes``."""

    import io

    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    for spec in slides:
        layout = presentation.slide_layouts[5 if "title" in spec else 6]
        slide = presentation.slides.add_slide(layout)
        if "title" in spec:
            slide.shapes.title.text = str(spec["title"])
        if "body" in spec:
            box = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(6), Inches(2))
            box.text_frame.text = "\n".join(spec["body"])  # type: ignore[arg-type]
        if "table" in spec:
            rows: list[list[str]] = spec["table"]  # type: ignore[assignment]
            shape = slide.shapes.add_table(
                len(rows), len(rows[0]), Inches(1), Inches(4), Inches(6), Inches(1)
            )
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    shape.table.cell(r, c).text = value
        if "notes" in spec:
            slide.notes_slide.notes_text_frame.text = str(spec["notes"])
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def odt_bytes(body_xml: str, *, extra_members: dict[str, bytes] | None = None) -> bytes:
    """An OpenDocument text file whose ``office:text`` holds ``body_xml``."""

    import zipfile

    content = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        "<office:document-content"
        ' xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"'
        ' xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"'
        ' xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0"'
        ' office:version="1.3">'
        f"<office:body><office:text>{body_xml}</office:text></office:body>"
        "</office:document-content>"
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        archive.writestr("content.xml", content)
        for name, data in (extra_members or {}).items():
            archive.writestr(name, data)
    return buffer.getvalue()
