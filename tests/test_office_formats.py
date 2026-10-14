"""XLSX and PPTX parsing: what is indexed, how it is labelled, and what is refused."""

from __future__ import annotations

import datetime
import io
import zipfile
from collections.abc import Callable

import pytest

from api.chunking import HeuristicTokenCounter
from api.documents import (
    SUPPORTED_EXTENSIONS,
    DocumentParseError,
    EmptyDocumentError,
    chunk_sections,
    parse_document_bytes,
)
from tests.factories import make_test_settings
from tests.samples import pptx_bytes, xlsx_bytes


def _rewrite_member(content: bytes, suffix: str, edit: Callable[[bytes], bytes]) -> bytes:
    """``content`` with the first member whose name ends in ``suffix`` passed through
    ``edit``."""

    source = zipfile.ZipFile(io.BytesIO(content))
    buffer = io.BytesIO()
    done = False
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info)
            if not done and info.filename.endswith(suffix):
                data = edit(data)
                done = True
            target.writestr(info, data)
    assert done
    return buffer.getvalue()


def _with_dtd(data: bytes) -> bytes:
    """A DTD right before the root element (after an XML declaration, if any)."""

    end = data.index(b"?>") + 2 if data.startswith(b"<?xml") else 0
    return data[:end] + b'<!DOCTYPE x [<!ENTITY a "aaaaaaaaaa">]>' + data[end:]


def test_both_formats_are_supported() -> None:
    assert {".xlsx", ".pptx"} <= SUPPORTED_EXTENSIONS


def test_xlsx_sheets_are_pages_with_spreadsheet_row_labels() -> None:
    content = xlsx_bytes(
        {
            "Salaries": [
                ["Name", "Role", "Start", "Bonus", "Remote"],
                ["Ada", "CEO", datetime.datetime(2020, 1, 5), 3.0, True],
                [],
                ["Grace", None, datetime.datetime(2021, 2, 1, 9, 30), 2.5, False],
            ],
            "Notes": [["Only a header"]],
            "Empty": [],
        }
    )

    sections = parse_document_bytes("pay.xlsx", content)

    assert [(s.page, s.section) for s in sections] == [
        (1, "Salaries · Rows 2-4"),
        (2, "Notes · Header"),
    ]
    assert sections[0].text == (
        "Name: Ada; Role: CEO; Start: 2020-01-05; Bonus: 3; Remote: TRUE\n"
        "Name: Grace; Start: 2021-02-01 09:30:00; Bonus: 2.5; Remote: FALSE"
    )
    assert sections[1].text == "Only a header"


def test_xlsx_rows_are_grouped_by_csv_rows_per_section() -> None:
    rows: list[list[object]] = [["n"], *[[index] for index in range(1, 6)]]
    settings = make_test_settings(csv_rows_per_section=2)

    sections = parse_document_bytes("n.xlsx", xlsx_bytes({"S": rows}), settings)

    assert [s.section for s in sections] == ["S · Rows 2-3", "S · Rows 4-5", "S · Row 6"]


def test_short_sheets_and_slides_stay_their_own_chunks() -> None:
    """Folding a short slide into the next one cited its text as the wrong slide."""

    settings = make_test_settings(chunk_size_tokens=200, chunk_overlap_tokens=10)
    deck = pptx_bytes([{"title": "One", "body": ["Tiny."]}, {"title": "Two", "body": ["Also."]}])
    book = xlsx_bytes({"A": [["k"], ["x"]], "B": [["k"], ["y"]]})

    for name, content in (("deck.pptx", deck), ("book.xlsx", book)):
        chunks = chunk_sections(
            parse_document_bytes(name, content, settings), settings, HeuristicTokenCounter()
        )
        assert [chunk.page for chunk in chunks] == [1, 2]


def test_xlsx_over_the_cell_limit_is_refused() -> None:
    content = xlsx_bytes({"S": [[1, 2, 3]] * 10})

    with pytest.raises(DocumentParseError, match="MAX_SPREADSHEET_CELLS"):
        parse_document_bytes("big.xlsx", content, make_test_settings(max_spreadsheet_cells=20))
    assert parse_document_bytes("big.xlsx", content, make_test_settings(max_spreadsheet_cells=40))


def test_xlsx_does_not_pad_rows_to_a_declared_dimension() -> None:
    """A sheet claiming A1:XFD1048576 would otherwise be padded to 17 billion cells."""

    def widen(data: bytes) -> bytes:
        widened = data.replace(b'<dimension ref="A1:A2"/>', b'<dimension ref="A1:XFD1048576"/>')
        assert widened != data
        return widened

    content = _rewrite_member(xlsx_bytes({"S": [["k"], ["v"]]}), "worksheets/sheet1.xml", widen)

    sections = parse_document_bytes("wide.xlsx", content, make_test_settings())

    assert sections[0].text == "k: v"


@pytest.mark.parametrize(
    ("name", "content", "suffix"),
    [
        ("book.xlsx", xlsx_bytes({"S": [["k"], ["v"]]}), "workbook.xml"),
        ("deck.pptx", pptx_bytes([{"title": "T", "body": ["b"]}]), "slide1.xml"),
    ],
)
def test_office_files_with_a_dtd_are_refused(name: str, content: bytes, suffix: str) -> None:
    hostile = _rewrite_member(content, suffix, _with_dtd)

    with pytest.raises(DocumentParseError, match="DTD"):
        parse_document_bytes(name, hostile)


@pytest.mark.parametrize("name", ["book.xlsx", "deck.pptx"])
def test_office_files_past_the_expanded_size_limit_are_refused(
    name: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    import api.parsers.common as common

    content = (
        xlsx_bytes({"S": [["k"], ["v"]]})
        if name.endswith(".xlsx")
        else pptx_bytes([{"title": "T", "body": ["b"]}])
    )
    module = "api.parsers.xlsx" if name.endswith(".xlsx") else "api.parsers.pptx"
    constant = "XLSX_MAX_EXPANDED_BYTES" if name.endswith(".xlsx") else "PPTX_MAX_EXPANDED_BYTES"
    monkeypatch.setattr(f"{module}.{constant}", 1024)
    assert common.ZIP_MAX_EXPANDED_BYTES > 1024

    with pytest.raises(DocumentParseError, match="MiB limit"):
        parse_document_bytes(name, content)


@pytest.mark.parametrize("name", ["junk.xlsx", "junk.pptx"])
def test_a_file_that_is_not_an_office_zip_is_a_parse_error(name: str) -> None:
    with pytest.raises(DocumentParseError):
        parse_document_bytes(name, b"not a zip at all")


def test_pptx_slides_are_pages_titled_by_their_slide() -> None:
    content = pptx_bytes(
        [
            {"title": "Roadmap", "body": ["Ship hybrid search.", "Add spans."]},
            {"title": "Numbers", "table": [["Quarter", "Users"], ["Q1", "40"], ["Q2", ""]]},
            {"body": ["A slide without a title."], "notes": "Say the date."},
            {"title": "Section break"},
            {},
        ]
    )

    sections = parse_document_bytes("deck.pptx", content)

    assert [(s.page, s.section) for s in sections] == [
        (1, "Roadmap"),
        (2, "Numbers"),
        (3, "Slide 3"),
        (4, "Section break"),
    ]
    # The title is the section label, not repeated as body text.
    assert sections[0].text == "Ship hybrid search.\nAdd spans."
    assert sections[1].text == "Quarter: Q1; Users: 40\nQuarter: Q2"
    assert sections[2].text == "A slide without a title.\nNotes: Say the date."
    assert sections[3].text == "Section break"


def test_pptx_reads_text_inside_grouped_shapes() -> None:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    box = group.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    box.text_frame.text = "Inside a group."
    buffer = io.BytesIO()
    presentation.save(buffer)

    [section] = parse_document_bytes("grouped.pptx", buffer.getvalue())

    assert section.text == "Inside a group."


def test_empty_workbooks_and_decks_are_empty_documents() -> None:
    with pytest.raises(EmptyDocumentError):
        parse_document_bytes("empty.xlsx", xlsx_bytes({"S": [[None, ""]]}))
    with pytest.raises(EmptyDocumentError):
        parse_document_bytes("empty.pptx", pptx_bytes([{}]))


def test_a_dtd_in_a_utf16_member_is_refused_too() -> None:
    """The scan matched only ASCII bytes; an XML parser reads UTF-16 from its BOM."""

    def utf16_with_dtd(data: bytes) -> bytes:
        text = data.decode("utf-8")
        if text.startswith("<?xml"):
            text = text[text.index("?>") + 2 :]
        hostile = '<?xml version="1.0" encoding="UTF-16"?><!DOCTYPE x [<!ENTITY a "a">]>' + text
        return hostile.encode("utf-16")

    content = _rewrite_member(xlsx_bytes({"S": [["k"], ["v"]]}), "workbook.xml", utf16_with_dtd)

    with pytest.raises(DocumentParseError, match="DTD"):
        parse_document_bytes("book.xlsx", content)


def test_declares_dtd_reads_every_xml_encoding() -> None:
    from api.parsers.common import declares_dtd

    for encoding in ("utf-8", "utf-16-le", "utf-16-be", "utf-32-le", "utf-32-be"):
        assert declares_dtd('<?xml version="1.0"?><!DOCTYPE x>'.encode(encoding)), encoding
    assert not declares_dtd("<root>no doctype</root>".encode("utf-16"))
