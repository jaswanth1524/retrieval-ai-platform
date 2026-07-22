from __future__ import annotations

from io import BytesIO
from typing import Any

import pytest

import api.documents as documents
from api.documents import (
    EmptyDocumentError,
    _table_to_markdown,
    parse_document_bytes,
)
from api.settings import AppSettings


def make_settings(**overrides: Any) -> AppSettings:
    return AppSettings(_env_file=None, **overrides)  # type: ignore[call-arg]


def _docx_bytes() -> bytes:
    from docx import Document

    document = Document()
    document.add_heading("Introduction", level=1)
    document.add_paragraph("DocRAG is a document question-answering system.")
    document.add_heading("Setup", level=2)
    document.add_paragraph("Run docker compose up to start.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Key"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Port"
    table.cell(1, 1).text = "8000"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_parse_docx_uses_headings_as_sections_and_includes_tables() -> None:
    sections = parse_document_bytes("guide.docx", _docx_bytes())

    labels = [section.section for section in sections]
    assert "Introduction" in labels
    assert "Setup" in labels
    assert all(section.page == 1 for section in sections)
    assert all(section.boundary == "semantic" for section in sections)
    combined = "\n".join(section.text for section in sections)
    assert "Port | 8000" in combined  # table rendered into text


def test_parse_html_strips_scripts_and_splits_on_headings() -> None:
    html = (
        "<html><head><title>Doc</title><style>.x{}</style></head>"
        "<body><script>evil()</script>"
        "<h1>Overview</h1><p>DocRAG does retrieval.</p>"
        "<h2>Details</h2><p>Hybrid search with reranking.</p></body></html>"
    )
    sections = parse_document_bytes("page.html", html.encode("utf-8"))

    labels = [section.section for section in sections]
    assert "Overview" in labels
    assert "Details" in labels
    combined = " ".join(section.text for section in sections)
    assert "evil()" not in combined  # script stripped
    assert "DocRAG does retrieval." in combined


def test_parse_csv_renders_rows_and_groups_into_sections() -> None:
    csv_text = "name,port\napi,8000\nqdrant,6333\nui,5173\n"
    settings = make_settings(csv_rows_per_section=2)

    sections = parse_document_bytes("services.csv", csv_text.encode("utf-8"), settings)

    # 3 data rows, grouped 2-per-section -> 2 sections.
    assert len(sections) == 2
    assert sections[0].section == "Rows 2-3"
    assert "name: api; port: 8000" in sections[0].text
    assert sections[1].section == "Rows 4-4"


def test_parse_csv_header_only_yields_header_section() -> None:
    sections = parse_document_bytes("empty.csv", b"name,port\n")
    assert len(sections) == 1
    assert sections[0].section == "Header"


def test_parse_pdf_appends_extracted_tables(monkeypatch: pytest.MonkeyPatch) -> None:
    class FakePage:
        def extract_text(self) -> str:
            return "Results\nThe measurements are below."

    class FakeReader:
        def __init__(self, stream: object) -> None:
            self.pages = [FakePage()]

    monkeypatch.setattr(documents, "PdfReader", FakeReader)
    monkeypatch.setattr(
        documents,
        "_extract_pdf_tables",
        lambda content: {1: ["| metric | value |\n| --- | --- |\n| latency | 5ms |"]},
    )

    sections = parse_document_bytes("report.pdf", b"not a real pdf")

    assert len(sections) == 1
    assert "[Table]" in sections[0].text
    assert "latency | 5ms" in sections[0].text


def test_parse_pdf_without_text_and_no_ocr_extra_points_at_install(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    class EmptyPage:
        def extract_text(self) -> str:
            return ""

    class FakeReader:
        def __init__(self, stream: object) -> None:
            self.pages = [EmptyPage()]

    monkeypatch.setattr(documents, "PdfReader", FakeReader)
    monkeypatch.setattr(documents, "_extract_pdf_tables", lambda content: {})

    # The rapidocr extra is not installed in the dev environment, so OCR import fails
    # and the error must point the user at the install command.
    with pytest.raises(EmptyDocumentError, match="uv sync --extra ocr"):
        parse_document_bytes("scanned.pdf", b"not a real pdf")


def test_parse_pdf_table_extraction_failure_degrades_to_text(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    class FakePage:
        def extract_text(self) -> str:
            return "Body text only."

    class FakeReader:
        def __init__(self, stream: object) -> None:
            self.pages = [FakePage()]

    monkeypatch.setattr(documents, "PdfReader", FakeReader)
    # _extract_pdf_tables opening the fake bytes fails internally and returns {} (a
    # logged, swallowed failure) — the page text must still parse cleanly.
    sections = parse_document_bytes("doc.pdf", b"not a real pdf")
    assert sections[0].text.startswith("Body text only.")


def test_table_to_markdown_handles_ragged_rows_and_none_cells() -> None:
    # Real pdfplumber output: ragged widths, None cells, embedded newlines.
    rows = [
        ["Metric", "Value", None],
        ["Latency", "5\nms"],
        [None, None, None],  # all-empty row is dropped
        ["Throughput", "100", "rps"],
    ]
    markdown = _table_to_markdown(rows)
    lines = markdown.splitlines()

    # Header + separator + two data rows (the all-None row is skipped).
    assert len(lines) == 4
    assert lines[0] == "| Metric | Value |  |"
    assert lines[1] == "| --- | --- | --- |"
    assert "5 ms" in lines[2]  # newline collapsed to a space
    assert lines[3] == "| Throughput | 100 | rps |"


def test_table_to_markdown_empty_returns_empty_string() -> None:
    assert _table_to_markdown([[None, None]]) == ""
    assert _table_to_markdown([]) == ""
